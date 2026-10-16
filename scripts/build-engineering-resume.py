from pathlib import Path
import json
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
import argparse
parser=argparse.ArgumentParser(description='Build the canonical two-page engineering resume DOCX')
parser.add_argument('--output-dir',required=True)
args=parser.parse_args()
ROOT=Path(__file__).resolve().parents[1]
WORK=ROOT/'work/mcdean-finalization'
OUT=Path(args.output_dir).resolve()
OUT.mkdir(parents=True,exist_ok=True)
WORK.mkdir(parents=True,exist_ok=True)
BASE='https://fatgezimb.github.io/fatgezim-portfolio/'
content=json.loads((ROOT/'app/content/engineering-resume.json').read_text())
current={'identity':{'email':content['email']}}
(WORK/'candidate-content.json').write_text(json.dumps(content,indent=2)+'\n')
doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=Inches(.58);sec.bottom_margin=Inches(.58)
sec.left_margin=Inches(.65);sec.right_margin=Inches(.65)
sec.header_distance=Inches(.22);sec.footer_distance=Inches(.24)
normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(11);normal.font.color.rgb=RGBColor(0,0,0)
normal.paragraph_format.line_spacing=Pt(14.2);normal.paragraph_format.space_after=Pt(3)
for name in ['Title','Subtitle','Heading 1','Heading 2','List Bullet']:
 s=doc.styles[name];s.font.name='Arial';s.font.color.rgb=RGBColor(0,0,0);s.font.italic=False
 for node in list(s.element.xpath('./w:pPr/w:pBdr')):node.getparent().remove(node)
for name,size in [('Title',25),('Subtitle',11),('Heading 1',11.5),('Heading 2',11)]:
 s=doc.styles[name];s.font.size=Pt(size);s.font.bold=True
 s.paragraph_format.keep_with_next=True
 s.paragraph_format.space_after=Pt(3)
 s.paragraph_format.space_before=Pt(9 if name=='Heading 1' else 0)
 s.paragraph_format.line_spacing=1.08
bul=doc.styles['List Bullet'];bul.font.size=Pt(11);bul.paragraph_format.left_indent=Inches(.13);bul.paragraph_format.first_line_indent=Inches(-.13);bul.paragraph_format.space_after=Pt(3);bul.paragraph_format.line_spacing=Pt(14.2)
doc.core_properties.title='Fatgezim Zim Bela Engineering Resume';doc.core_properties.subject='Software engineering data engineering and technical leadership';doc.core_properties.author='Fatgezim Bela';doc.core_properties.last_modified_by='Fatgezim Bela';doc.core_properties.keywords='Software Engineer, Data Engineering, Technical Leadership'

def para(text='',style=None,bold=None,after=None):
 p=doc.add_paragraph(style=style);run=p.add_run(text)
 if bold is not None:run.bold=bold
 if after is not None:p.paragraph_format.space_after=Pt(after)
 return p

def link(p,label,url,size=9.5):
 h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
 run=OxmlElement('w:r');pr=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'000000');pr.append(color)
 fonts=OxmlElement('w:rFonts');fonts.set(qn('w:ascii'),'Arial');fonts.set(qn('w:hAnsi'),'Arial');pr.append(fonts)
 sz=OxmlElement('w:sz');sz.set(qn('w:val'),str(round(size*2)));pr.append(sz)
 run.append(pr);t=OxmlElement('w:t');t.text=label;run.append(t);h.append(run);p._p.append(h)

def section(title):return para(title,'Heading 1')
def item(title,org,dates):
 p=para(style='Heading 2');p.paragraph_format.space_before=Pt(6);p.paragraph_format.tab_stops.add_tab_stop(Inches(7.2),WD_TAB_ALIGNMENT.RIGHT)
 p.add_run(title).bold=True;p.add_run('\t'+dates).font.size=Pt(9.5)
 p=para(org,after=3);p.paragraph_format.keep_with_next=True

def bullet(text):return para(text,'List Bullet')

p=para(content['name'],'Title');p.paragraph_format.space_after=Pt(4)
para(content['headline'],'Subtitle',after=6)
p=doc.add_paragraph();p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=Pt(12)
contacts=[(current['identity']['email'],'mailto:'+current['identity']['email']),('github.com/Fatgezimb','https://github.com/Fatgezimb'),('linkedin.com/in/fatgezimzimbela','https://www.linkedin.com/in/fatgezimzimbela/')]
for i,(label,url) in enumerate(contacts):
 if i:p.add_run('  |  ').font.size=Pt(9.5)
 link(p,label,url)
p=doc.add_paragraph();p.paragraph_format.space_after=Pt(5);p.paragraph_format.line_spacing=Pt(12)
link(p,'fatgezimb.github.io/fatgezim-portfolio',BASE)
section('Professional Summary');para(content['summary'])
section('Engineering Skills')
for label,text in content['skills']:
 p=doc.add_paragraph();p.paragraph_format.space_after=Pt(3);p.add_run(label+'  ').bold=True;p.add_run(text)
section('Relevant Engineering and Clinical Experience')
for role in content['experience']:
 item(role['title'],role['organization'],role['dates'])
 for b in role['bullets']:bullet(b)
p=doc.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0);p.add_run().add_break(WD_BREAK.PAGE)
section('Selected Engineering Projects')
for project in content['projects']:
 p=para(project['name'],'Heading 2');p.paragraph_format.space_before=Pt(6)
 if project['link']:
  p.paragraph_format.tab_stops.add_tab_stop(Inches(7.2),WD_TAB_ALIGNMENT.RIGHT);p.add_run('\t');link(p,'View project',project['link'],9)
 p=para(project['meta'],after=4);p.paragraph_format.keep_with_next=True
 for run in p.runs:run.font.size=Pt(9.5)
 for b in project['bullets']:bullet(b)
section('Research')
p=para(content['research']['title']+' | '+content['research']['institution'],'Heading 2');p.paragraph_format.tab_stops.add_tab_stop(Inches(7.2),WD_TAB_ALIGNMENT.RIGHT);p.add_run('\t');link(p,'Poster',BASE+'media/fatgezim-bela-computational-neuroscience-poster.pdf',9)
para(content['research']['text'],after=3)
para('Methods: '+content['research']['methods'],after=2)
section('Education')
for e in content['education']:
 p=doc.add_paragraph();p.paragraph_format.space_after=Pt(3);p.add_run(e['credential']).bold=True;p.add_run(' | '+e['institution']+' | '+e['date'])
section('Professional Credentials')
for sentence in content['credentials'].split('. '):para(sentence if sentence.endswith('.') else sentence+'.',after=2)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT;footer.paragraph_format.space_after=Pt(0);footer.paragraph_format.line_spacing=1
footer.style.font.name='Arial';footer.style.font.size=Pt(8);footer.style.font.color.rgb=RGBColor.from_string('555555')
footer.paragraph_format.tab_stops.clear_all();footer.paragraph_format.tab_stops.add_tab_stop(Inches(7.2),WD_TAB_ALIGNMENT.RIGHT)
r=footer.add_run('Fatgezim Zim Bela  |  ');r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string('555555')
for value in ['PAGE','NUMPAGES']:
 if value=='NUMPAGES':footer.add_run(' / ').font.size=Pt(8)
 r=footer.add_run();r.font.size=Pt(8);f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),value);r._r.addnext(f)
settings=doc.settings.element
u=OxmlElement('w:updateFields');u.set(qn('w:val'),'true');settings.append(u)
path=OUT/'Fatgezim_Zim_Bela_Engineering_Resume.docx';doc.save(path)
print(json.dumps({'docx':str(path),'bodyFontPoints':11,'wordCount':sum(len(p.text.split()) for p in doc.paragraphs)},indent=2))
